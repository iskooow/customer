"""
Tests for views.
"""

import pytest
from datetime import date, timedelta
from decimal import Decimal

from django.test import TestCase
from django.urls import reverse
from django.utils import timezone

from accounts.models import User
from customers.models import Customer
from salespersons.models import SalesPerson
from documents.models import Document
from security_cheques.models import SecurityCheque


@pytest.mark.django_db
class TestCustomerViews(TestCase):
    def setUp(self):
        self.admin_user = User.objects.create_user(
            email='admin@example.com',
            username='admin',
            password='admin123',
            role=User.Role.ADMIN,
            is_staff=True,
            is_superuser=True,
        )
        
        self.sales_user = User.objects.create_user(
            email='sales@example.com',
            username='sales',
            password='sales123',
            role=User.Role.SALES_PERSON,
        )
        
        self.sales_person = SalesPerson.objects.create(
            name='John Doe',
            email='john@example.com',
            phone='+971 50 123 4567',
        )
        
        # Link sales user to sales person
        self.sales_user.sales_person = self.sales_person
        self.sales_user.save()
        
        self.customer = Customer.objects.create(
            company_name='Test Company LLC',
            trade_license_number='123456',
            trade_license_expiry=date.today() + timedelta(days=30),
            trn='100123456789003',
            sales_person=self.sales_person,
            created_by=self.admin_user,
            updated_by=self.admin_user,
        )
    
    def test_dashboard_access_admin(self):
        self.client.login(email='admin@example.com', password='admin123')
        response = self.client.get(reverse('dashboard'))
        self.assertEqual(response.status_code, 200)
    
    def test_dashboard_access_sales_person(self):
        self.client.login(email='sales@example.com', password='sales123')
        response = self.client.get(reverse('dashboard'))
        self.assertEqual(response.status_code, 200)
    
    def test_customer_list_access_admin(self):
        self.client.login(email='admin@example.com', password='admin123')
        response = self.client.get(reverse('customers:list'))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Test Company LLC')
    
    def test_customer_list_access_sales_person(self):
        self.client.login(email='sales@example.com', password='sales123')
        response = self.client.get(reverse('customers:list'))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Test Company LLC')
    
    def test_customer_list_redirect_anonymous(self):
        response = self.client.get(reverse('customers:list'))
        self.assertEqual(response.status_code, 302)  # Redirect to login

    def test_customer_list_htmx_returns_partial(self):
        """Regression: HTMX requests to the customer list must return only the
        table partial (no base.html wrapper) so #customer-table swaps cleanly
        without re-rendering the full page or duplicating the filters card."""
        self.client.login(email='admin@example.com', password='admin123')
        response = self.client.get(reverse('customers:list'), HTTP_HX_REQUEST='true')
        self.assertEqual(response.status_code, 200)
        content = response.content.decode()
        self.assertNotIn('<html', content)
        self.assertNotIn('sidebar-link', content)
        self.assertIn('customer-table-body', content)

    def test_customer_list_htmx_search_filters(self):
        """Regression: HTMX search must filter the rows rendered in the
        partial response, not the whole page."""
        Customer.objects.create(
            company_name='Alpha Search Co LLC',
            trade_license_expiry=date.today() + timedelta(days=30),
            sales_person=self.sales_person,
            created_by=self.admin_user,
            updated_by=self.admin_user,
        )
        Customer.objects.create(
            company_name='Beta Other Co LLC',
            trade_license_expiry=date.today() + timedelta(days=30),
            sales_person=self.sales_person,
            created_by=self.admin_user,
            updated_by=self.admin_user,
        )
        self.client.login(email='admin@example.com', password='admin123')
        response = self.client.get(
            reverse('customers:list') + '?search=Alpha', HTTP_HX_REQUEST='true'
        )
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Alpha Search Co LLC')
        self.assertNotContains(response, 'Beta Other Co LLC')

    def test_customer_list_pagination_preserves_filters(self):
        """Pagination links in the partial must retain the active filter
        querystring so changing pages does not reset the current filters."""
        for i in range(26):
            Customer.objects.create(
                company_name=f'Bulk Co {i:03d}',
                trade_license_expiry=date.today() + timedelta(days=30),
                sales_person=self.sales_person,
                created_by=self.admin_user,
                updated_by=self.admin_user,
            )
        self.client.login(email='admin@example.com', password='admin123')
        response = self.client.get(
            reverse('customers:list') + '?page=1&search=Bulk',
            HTTP_HX_REQUEST='true',
        )
        self.assertEqual(response.status_code, 200)
        content = response.content.decode()
        self.assertIn('?page=2&search=Bulk', content)
        self.assertIn('hx-target="#customer-table"', content)
    
    def test_customer_detail_access_admin(self):
        self.client.login(email='admin@example.com', password='admin123')
        response = self.client.get(reverse('customers:detail', kwargs={'pk': self.customer.pk}))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Test Company LLC')
    
    def test_customer_detail_access_sales_person_own(self):
        self.client.login(email='sales@example.com', password='sales123')
        response = self.client.get(reverse('customers:detail', kwargs={'pk': self.customer.pk}))
        self.assertEqual(response.status_code, 200)
    
    def test_customer_detail_access_sales_person_other(self):
        # Create another sales person and customer
        other_sp = SalesPerson.objects.create(
            name='Jane Smith',
            email='jane@example.com',
        )
        other_customer = Customer.objects.create(
            company_name='Other Company LLC',
            sales_person=other_sp,
            created_by=self.admin_user,
            updated_by=self.admin_user,
        )
        
        self.client.login(email='sales@example.com', password='sales123')
        response = self.client.get(reverse('customers:detail', kwargs={'pk': other_customer.pk}))
        self.assertEqual(response.status_code, 302)  # Redirect due to permission
    
    def test_add_customer_admin(self):
        self.client.login(email='admin@example.com', password='admin123')
        response = self.client.get(reverse('customers:add'))
        self.assertEqual(response.status_code, 200)
        
        # Test POST
        data = {
            'company_name': 'New Company LLC',
            'trade_license_number': '654321',
            'trade_license_expiry': (date.today() + timedelta(days=60)).strftime('%Y-%m-%d'),
            'trn': '100987654321003',
            'sales_person': self.sales_person.pk,
            'customer_status': 'active',
            'passport_number': 'A1234567',
            'passport_expiry': (date.today() + timedelta(days=365)).strftime('%Y-%m-%d'),
            'eid_number': '784-2000-1234567-1',
            'eid_expiry': (date.today() + timedelta(days=730)).strftime('%Y-%m-%d'),
        }
        response = self.client.post(reverse('customers:add'), data)
        self.assertEqual(response.status_code, 302)  # Redirect after success
        self.assertTrue(Customer.objects.filter(company_name='New Company LLC').exists())

    def test_add_customer_without_trn_or_trade_license(self):
        """Empty optional TRN/trade license fields must not crash the form.

        Regression for AttributeError: 'NoneType' object has no attribute
        'strip' in CustomerForm.clean_trn / clean_trade_license_number when a
        blank nullable CharField yields None in cleaned_data.
        """
        self.client.login(email='admin@example.com', password='admin123')
        data = {
            'company_name': 'No Trn Company LLC',
            'trade_license_number': '',
            'trn': '',
            'sales_person': self.sales_person.pk,
            'customer_status': 'active',
        }
        response = self.client.post(reverse('customers:add'), data)
        self.assertEqual(response.status_code, 302)  # Redirect after success
        customer = Customer.objects.get(company_name='No Trn Company LLC')
        self.assertIsNone(customer.trn)
        self.assertIsNone(customer.trade_license_number)

    def test_add_customer_whitespace_trn_stored_as_null(self):
        """A whitespace-only TRN is stripped to empty and stored as NULL."""
        self.client.login(email='admin@example.com', password='admin123')
        data = {
            'company_name': 'Whitespace Trn Company LLC',
            'trade_license_number': '   ',
            'trn': '   ',
            'sales_person': self.sales_person.pk,
            'customer_status': 'active',
        }
        response = self.client.post(reverse('customers:add'), data)
        self.assertEqual(response.status_code, 302)
        customer = Customer.objects.get(company_name='Whitespace Trn Company LLC')
        self.assertIsNone(customer.trn)
        self.assertIsNone(customer.trade_license_number)

    def test_add_customer_sales_person(self):
        self.client.login(email='sales@example.com', password='sales123')
        response = self.client.get(reverse('customers:add'))
        self.assertEqual(response.status_code, 200)
        
        data = {
            'company_name': 'Sales Person Company LLC',
            'trade_license_number': '111222',
            'trade_license_expiry': (date.today() + timedelta(days=60)).strftime('%Y-%m-%d'),
            'trn': '100111222333444',
            'sales_person': self.sales_person.pk,
            'customer_status': 'active',
        }
        response = self.client.post(reverse('customers:add'), data)
        self.assertEqual(response.status_code, 302)
        customer = Customer.objects.get(company_name='Sales Person Company LLC')
        self.assertEqual(customer.sales_person, self.sales_person)
    
    def test_edit_customer_admin(self):
        self.client.login(email='admin@example.com', password='admin123')
        response = self.client.get(reverse('customers:edit', kwargs={'pk': self.customer.pk}))
        self.assertEqual(response.status_code, 200)
        
        data = {
            'company_name': 'Updated Company LLC',
            'trade_license_number': '123456',
            'trade_license_expiry': (date.today() + timedelta(days=30)).strftime('%Y-%m-%d'),
            'trn': '100123456789003',
            'sales_person': self.sales_person.pk,
            'customer_status': 'active',
        }
        response = self.client.post(reverse('customers:edit', kwargs={'pk': self.customer.pk}), data)
        self.assertEqual(response.status_code, 302)
        self.customer.refresh_from_db()
        self.assertEqual(self.customer.company_name, 'Updated Company LLC')
    
    def test_edit_customer_reconciles_expiry_from_documents(self):
        """
        Regression: editing the customer form must not regress the passport
        expiry back to a stale value when an uploaded passport document has a
        newer (future) expiry. Saving the form re-syncs from the documents.
        """
        future_expiry = date.today() + timedelta(days=365)
        # A renewed passport document exists with a future expiry...
        Document.objects.create(
            customer=self.customer,
            document_type=Document.DocumentType.PASSPORT,
            file_name='renewed_passport.pdf',
            file_size=1024,
            expiry_date=future_expiry,
            uploaded_by=self.admin_user,
        )
        # ...but the customer record still holds the stale, expired value.
        self.customer.passport_expiry = date.today() - timedelta(days=30)
        self.customer.save(update_fields=['passport_expiry', 'updated_at'])

        self.client.login(email='admin@example.com', password='admin123')
        data = {
            'company_name': self.customer.company_name,
            'trade_license_number': self.customer.trade_license_number,
            'trade_license_expiry': self.customer.trade_license_expiry.strftime('%Y-%m-%d') if self.customer.trade_license_expiry else '',
            'trn': self.customer.trn,
            'sales_person': self.sales_person.pk,
            'customer_status': 'active',
        }
        response = self.client.post(reverse('customers:edit', kwargs={'pk': self.customer.pk}), data)
        self.assertEqual(response.status_code, 302)

        self.customer.refresh_from_db()
        # The form save must reconcile the expiry up to the document's future date.
        self.assertEqual(self.customer.passport_expiry, future_expiry)
    
    def test_delete_customer_admin_only(self):
        self.client.login(email='sales@example.com', password='sales123')
        response = self.client.post(reverse('customers:delete', kwargs={'pk': self.customer.pk}))
        self.assertEqual(response.status_code, 302)  # Redirect with error
        self.assertTrue(Customer.objects.filter(pk=self.customer.pk).exists())
        
        self.client.login(email='admin@example.com', password='admin123')
        response = self.client.post(reverse('customers:delete', kwargs={'pk': self.customer.pk}))
        self.assertEqual(response.status_code, 302)  # Redirect to list
        self.assertFalse(Customer.objects.filter(pk=self.customer.pk).exists())
    
    def test_customer_search(self):
        self.client.login(email='admin@example.com', password='admin123')
        response = self.client.get(reverse('customers:list'), {'search': 'Test Company'})
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Test Company LLC')
        
        response = self.client.get(reverse('customers:list'), {'search': 'NonExistent'})
        self.assertEqual(response.status_code, 200)
        self.assertNotContains(response, 'Test Company LLC')
    
    def test_customer_filter_by_sales_person(self):
        self.client.login(email='admin@example.com', password='admin123')
        response = self.client.get(reverse('customers:list'), {'sales_person': self.sales_person.pk})
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Test Company LLC')
    
    def test_customer_filter_by_document_status(self):
        # Create customer with expired document
        expired_customer = Customer.objects.create(
            company_name='Expired Company LLC',
            trade_license_expiry=date.today() - timedelta(days=10),
            sales_person=self.sales_person,
            created_by=self.admin_user,
            updated_by=self.admin_user,
        )
        
        self.client.login(email='admin@example.com', password='admin123')
        response = self.client.get(reverse('customers:list'), {'document_status': 'expired'})
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Expired Company LLC')
        self.assertNotContains(response, 'Test Company LLC')


@pytest.mark.django_db
class TestDocumentViews(TestCase):
    def setUp(self):
        self.admin_user = User.objects.create_user(
            email='admin@example.com',
            username='admin',
            password='admin123',
            role=User.Role.ADMIN,
        )
        
        self.sales_person = SalesPerson.objects.create(
            name='John Doe',
            email='john@example.com',
        )
        
        self.customer = Customer.objects.create(
            company_name='Test Company LLC',
            sales_person=self.sales_person,
            created_by=self.admin_user,
            updated_by=self.admin_user,
        )
    
    def test_document_upload(self):
        self.client.login(email='admin@example.com', password='admin123')
        response = self.client.get(reverse('documents:upload', kwargs={'customer_pk': self.customer.pk}))
        self.assertEqual(response.status_code, 200)
    
    def test_document_list(self):
        self.client.login(email='admin@example.com', password='admin123')
        response = self.client.get(reverse('documents:list'))
        self.assertEqual(response.status_code, 200)

    def test_document_view_and_download_admin_no_500(self):
        """Regression: DocumentAccessMixin must not call self.get_object()
        on plain View subclasses (DocumentViewView / DocumentDownloadView)."""
        doc = Document.objects.create(
            customer=self.customer,
            document_type=Document.DocumentType.TRADE_LICENSE,
            file_name='missing.pdf',
            file_size=1000,
            uploaded_by=self.admin_user,
        )
        self.client.login(email='admin@example.com', password='admin123')
        # With no actual file on disk both views redirect to the customer detail
        # page ("File not found") instead of raising AttributeError.
        for url_name in ('documents:view', 'documents:download'):
            response = self.client.get(reverse(url_name, kwargs={'pk': doc.pk}))
            self.assertEqual(response.status_code, 302, url_name)
            self.assertEqual(
                response['Location'],
                reverse('customers:detail', kwargs={'pk': self.customer.pk}),
            )

    def test_document_view_sales_person_own_customer(self):
        """Sales person can access documents of their own customer (no 500)."""
        sales_user = User.objects.create_user(
            email='sales@example.com',
            username='sales',
            password='sales123',
            role=User.Role.SALES_PERSON,
            sales_person=self.sales_person,
        )
        doc = Document.objects.create(
            customer=self.customer,
            document_type=Document.DocumentType.PASSPORT,
            file_name='missing.pdf',
            file_size=1000,
            uploaded_by=sales_user,
        )
        self.client.login(email='sales@example.com', password='sales123')
        response = self.client.get(reverse('documents:view', kwargs={'pk': doc.pk}))
        self.assertEqual(response.status_code, 302)


@pytest.mark.django_db
class TestSalesPersonViews(TestCase):
    def setUp(self):
        self.admin_user = User.objects.create_user(
            email='admin@example.com',
            username='admin',
            password='admin123',
            role=User.Role.ADMIN,
        )
        
        self.sales_person = SalesPerson.objects.create(
            name='John Doe',
            email='john@example.com',
        )
    
    def test_salesperson_list_admin(self):
        self.client.login(email='admin@example.com', password='admin123')
        response = self.client.get(reverse('salespersons:list'))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'John Doe')
    
    def test_salesperson_list_sales_person_denied(self):
        sales_user = User.objects.create_user(
            email='sales@example.com',
            username='sales',
            password='sales123',
            role=User.Role.SALES_PERSON,
        )
        self.client.login(email='sales@example.com', password='sales123')
        response = self.client.get(reverse('salespersons:list'))
        self.assertEqual(response.status_code, 302)  # Redirect with error

    def test_salesperson_list_htmx_returns_partial(self):
        """Regression: HTMX requests to the salesperson list must return only the
        table partial (no base.html wrapper) so #salespersons-table swaps cleanly
        without re-rendering the whole page or duplicating the container."""
        self.client.login(email='admin@example.com', password='admin123')
        response = self.client.get(
            reverse('salespersons:list'), HTTP_HX_REQUEST='true'
        )
        self.assertEqual(response.status_code, 200)
        content = response.content.decode()
        self.assertNotIn('<html', content)
        self.assertNotIn('sidebar-link', content)
        self.assertIn('table-container', content)

    def test_salesperson_list_htmx_search_filters(self):
        """Regression: the search filter must reach the queryset on an HTMX
        request so the partial only renders matching sales persons."""
        SalesPerson.objects.create(
            name='Alpha Finder',
            email='alpha@example.com',
            active=True,
        )
        SalesPerson.objects.create(
            name='Beta Other',
            email='beta@example.com',
            active=True,
        )
        self.client.login(email='admin@example.com', password='admin123')
        self.assertContains(
            self.client.get(reverse('salespersons:list') + '?search=Alpha'),
            'Alpha Finder',
        )
        response = self.client.get(
            reverse('salespersons:list') + '?search=Alpha', HTTP_HX_REQUEST='true'
        )
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Alpha Finder')
        self.assertNotContains(response, 'Beta Other')

    def test_salesperson_list_htmx_status_filter(self):
        """Regression: the status filter must reach the queryset on an HTMX
        request so the partial only renders sales persons matching the status."""
        SalesPerson.objects.create(
            name='Active Person',
            email='active-p@example.com',
            active=True,
        )
        SalesPerson.objects.create(
            name='Inactive Person',
            email='inactive-p@example.com',
            active=False,
        )
        self.client.login(email='admin@example.com', password='admin123')
        response = self.client.get(
            reverse('salespersons:list') + '?status=inactive', HTTP_HX_REQUEST='true'
        )
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Inactive Person')
        self.assertNotContains(response, 'Active Person')



@pytest.mark.django_db
class TestSecurityChequeViews(TestCase):
    def setUp(self):
        self.admin_user = User.objects.create_user(
            email='admin@example.com',
            username='admin',
            password='admin123',
            role=User.Role.ADMIN,
        )
        
        self.sales_person = SalesPerson.objects.create(
            name='John Doe',
            email='john@example.com',
        )
        
        self.customer = Customer.objects.create(
            company_name='Test Company LLC',
            sales_person=self.sales_person,
            created_by=self.admin_user,
            updated_by=self.admin_user,
        )
    
    def test_security_cheque_list(self):
        self.client.login(email='admin@example.com', password='admin123')
        response = self.client.get(reverse('security_cheques:list'))
        self.assertEqual(response.status_code, 200)

    def test_security_cheque_list_htmx_returns_partial(self):
        """Regression: HTMX requests to the cheque list must return only the
        table partial (no base.html wrapper) so #cheques-table swaps cleanly
        without re-rendering the whole page or duplicating the container."""
        self.client.login(email='admin@example.com', password='admin123')
        response = self.client.get(
            reverse('security_cheques:list'), HTTP_HX_REQUEST='true'
        )
        self.assertEqual(response.status_code, 200)
        content = response.content.decode()
        self.assertNotIn('<html', content)
        self.assertNotIn('sidebar-link', content)
        self.assertIn('table-container', content)

    def test_security_cheque_list_htmx_search_filters(self):
        """Regression: the search filter must reach the queryset on an HTMX
        request (reports view previously ignored the search param entirely)."""
        SecurityCheque.objects.create(
            customer=self.customer,
            status='received',
            cheque_number='CHQ-FIND-ME',
            amount=Decimal('1000.00'),
            created_by=self.admin_user,
            updated_by=self.admin_user,
        )
        SecurityCheque.objects.create(
            customer=self.customer,
            status='pending',
            cheque_number='CHQ-IGNORED',
            amount=Decimal('2000.00'),
            created_by=self.admin_user,
            updated_by=self.admin_user,
        )
        self.client.login(email='admin@example.com', password='admin123')

        # Full list shows both
        response = self.client.get(reverse('security_cheques:list'))
        self.assertContains(response, 'CHQ-FIND-ME')
        self.assertContains(response, 'CHQ-IGNORED')

        # HTMX search narrows the results
        response = self.client.get(
            reverse('security_cheques:list') + '?search=FIND-ME',
            HTTP_HX_REQUEST='true',
        )
        self.assertContains(response, 'CHQ-FIND-ME')
        self.assertNotContains(response, 'CHQ-IGNORED')

    def test_security_cheque_list_pagination_preserves_filters(self):
        """Pagination links must retain the active search/status querystring
        so changing pages does not reset the current filters."""
        for i in range(26):
            SecurityCheque.objects.create(
                customer=self.customer,
                status='received',
                cheque_number=f'CHQ-{i:03d}',
                amount=Decimal('100.00'),
                created_by=self.admin_user,
                updated_by=self.admin_user,
            )
        self.client.login(email='admin@example.com', password='admin123')
        response = self.client.get(
            reverse('security_cheques:list') + '?page=1&search=CHQ&status=received'
        )
        self.assertEqual(response.status_code, 200)
        content = response.content.decode()
        self.assertIn('?page=2&search=CHQ&status=received', content)
        self.assertIn('hx-target="#cheques-table"', content)
    
    def test_security_cheque_add(self):
        self.client.login(email='admin@example.com', password='admin123')
        response = self.client.get(reverse('security_cheques:add', kwargs={'customer_pk': self.customer.pk}))
        self.assertEqual(response.status_code, 200)
        
        data = {
            'status': 'received',
            'cheque_number': 'CHQ123456',
            'amount': '10000.00',
            'cheque_date': date.today().strftime('%Y-%m-%d'),
        }
        response = self.client.post(reverse('security_cheques:add', kwargs={'customer_pk': self.customer.pk}), data)
        self.assertEqual(response.status_code, 302)
        self.assertTrue(SecurityCheque.objects.filter(cheque_number='CHQ123456').exists())

    def test_security_cheque_download_admin_no_500(self):
        """Regression: ChequeAccessMixin must not call self.get_object()
        on plain View subclasses (SecurityChequeDownloadView)."""
        cheque = SecurityCheque.objects.create(
            customer=self.customer,
            status=SecurityCheque.Status.RECEIVED,
            cheque_number='CHQ789',
            amount=Decimal('5000.00'),
            cheque_date=date.today(),
            created_by=self.admin_user,
            updated_by=self.admin_user,
        )
        self.client.login(email='admin@example.com', password='admin123')
        # No actual file on disk -> view redirects to customer detail ("File not
        # found") instead of raising AttributeError.
        response = self.client.get(reverse('security_cheques:download', kwargs={'pk': cheque.pk}))
        self.assertEqual(response.status_code, 302)
        self.assertEqual(
            response['Location'],
            reverse('customers:detail', kwargs={'pk': self.customer.pk}),
        )


@pytest.mark.django_db
class TestExportViews(TestCase):
    def setUp(self):
        self.admin_user = User.objects.create_user(
            email='admin@example.com',
            username='admin',
            password='admin123',
            role=User.Role.ADMIN,
        )
        
        self.sales_person = SalesPerson.objects.create(
            name='John Doe',
            email='john@example.com',
        )
        
        self.customer = Customer.objects.create(
            company_name='Test Company LLC',
            trade_license_number='123456',
            trade_license_expiry=date.today() + timedelta(days=30),
            trn='100123456789003',
            sales_person=self.sales_person,
            created_by=self.admin_user,
            updated_by=self.admin_user,
        )
    
    def test_excel_export(self):
        self.client.login(email='admin@example.com', password='admin123')
        response = self.client.get(reverse('customers:export'))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response['Content-Type'],
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        self.assertIn('attachment', response['Content-Disposition'])
    
    def test_csv_export(self):
        self.client.login(email='admin@example.com', password='admin123')
        response = self.client.get(reverse('customers:export'), {'format': 'csv'})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response['Content-Type'], 'text/csv')
        self.assertIn('attachment', response['Content-Disposition'])


@pytest.mark.django_db
class TestReportViews(TestCase):
    def setUp(self):
        self.admin_user = User.objects.create_user(
            email='admin@example.com',
            username='admin',
            password='admin123',
            role=User.Role.ADMIN,
        )
        
        self.sales_person = SalesPerson.objects.create(
            name='John Doe',
            email='john@example.com',
        )
        
        self.customer = Customer.objects.create(
            company_name='Test Company LLC',
            trade_license_number='123456',
            trade_license_expiry=date.today() + timedelta(days=30),
            trn='100123456789003',
            sales_person=self.sales_person,
            created_by=self.admin_user,
            updated_by=self.admin_user,
        )
    
    def test_reports_dashboard(self):
        self.client.login(email='admin@example.com', password='admin123')
        response = self.client.get(reverse('dashboard'))
        self.assertEqual(response.status_code, 200)
    
    def test_customer_report(self):
        self.client.login(email='admin@example.com', password='admin123')
        response = self.client.get(reverse('reports:customer_report'))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Test Company LLC')
    
    def test_expired_documents_report(self):
        # Create expired customer
        expired_customer = Customer.objects.create(
            company_name='Expired Company LLC',
            trade_license_expiry=date.today() - timedelta(days=10),
            sales_person=self.sales_person,
            created_by=self.admin_user,
            updated_by=self.admin_user,
        )
        
        self.client.login(email='admin@example.com', password='admin123')
        response = self.client.get(reverse('reports:expired_documents'))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Expired Company LLC')
    
    def test_salesperson_report(self):
        self.client.login(email='admin@example.com', password='admin123')
        response = self.client.get(reverse('reports:salesperson_report'))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'John Doe')
    
    def test_security_cheque_report(self):
        self.client.login(email='admin@example.com', password='admin123')
        response = self.client.get(reverse('reports:cheque_report'))
        self.assertEqual(response.status_code, 200)

    def test_security_cheque_report_htmx_returns_partial(self):
        """Regression: HTMX requests to the cheque report must return only the
        table partial so #cheques-table swaps cleanly on filter/search."""
        self.client.login(email='admin@example.com', password='admin123')
        response = self.client.get(
            reverse('reports:cheque_report'), HTTP_HX_REQUEST='true'
        )
        self.assertEqual(response.status_code, 200)
        content = response.content.decode()
        self.assertNotIn('<html', content)
        self.assertNotIn('sidebar-link', content)
        self.assertIn('table-container', content)

    def test_security_cheque_report_search_filters(self):
        """Regression: the cheque report search input must actually filter the
        queryset (the view previously only supported the status filter)."""
        SecurityCheque.objects.create(
            customer=self.customer,
            status='received',
            cheque_number='CHQ-REPORT-ALPHA',
            amount=Decimal('1000.00'),
            created_by=self.admin_user,
            updated_by=self.admin_user,
        )
        SecurityCheque.objects.create(
            customer=self.customer,
            status='pending',
            cheque_number='CHQ-REPORT-BETA',
            amount=Decimal('2000.00'),
            created_by=self.admin_user,
            updated_by=self.admin_user,
        )
        self.client.login(email='admin@example.com', password='admin123')
        response = self.client.get(
            reverse('reports:cheque_report') + '?search=ALPHA',
            HTTP_HX_REQUEST='true',
        )
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'CHQ-REPORT-ALPHA')
        self.assertNotContains(response, 'CHQ-REPORT-BETA')

    def test_customer_report_excel_export(self):
        """Regression: exporting the customer report to Excel must not raise
        'Cannot convert ... to Excel' from a lazy translation proxy written to
        an openpyxl cell (the 'Summary' heading in add_summary)."""
        self.client.login(email='admin@example.com', password='admin123')
        # The reported failure URL included ?format=excel&report_type=salesperson
        response = self.client.get(
            reverse('reports:customer_report') + '?format=excel&report_type=salesperson'
        )
        self.assertEqual(response.status_code, 200)
        self.assertIn(
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            response['Content-Type'],
        )

    def _load_report_workbook(self, url):
        """GET the URL, load the returned xlsx into an openpyxl Workbook."""
        import io
        import openpyxl
        # Ensure we are authenticated before hitting the report endpoint.
        if not self.client.login(email='admin@example.com', password='admin123'):
            raise AssertionError('Could not log in test admin user')
        response = self.client.get(url)
        self.assertEqual(response.status_code, 200)
        self.assertIn(
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            response['Content-Type'],
        )
        return openpyxl.load_workbook(io.BytesIO(response.content)).active

    def test_customer_report_excel_has_kb_branding(self):
        """The redesigned report header carries the KB Remicon branding."""
        ws = self._load_report_workbook(
            reverse('reports:customer_report') + '?format=excel'
        )
        self.assertEqual(ws['A1'].value, 'KB REMICON')
        self.assertEqual(ws['A2'].value, 'CUSTOMER RECORD')
        self.assertEqual(ws['E1'].value, 'KB REMICON CUSTOMER RECORD')
        self.assertEqual(ws['E2'].value, 'CUSTOMER COMPLIANCE REPORT')

    def test_customer_report_excel_has_metadata(self):
        """The report shows who generated it and when."""
        ws = self._load_report_workbook(
            reverse('reports:customer_report') + '?format=excel'
        )
        self.assertTrue(ws['L1'].value.startswith('Report Type'))
        # Generated By falls back to the username (email) when no full name is set.
        # The admin test user uses admin@example.com as its USERNAME_FIELD value.
        self.assertEqual(ws['L2'].value, 'Generated By    : admin@example.com')
        self.assertTrue(ws['L3'].value.startswith('Generated On    :'))

    def test_customer_report_excel_has_summary(self):
        """The summary section shows total/active counts and date."""
        ws = self._load_report_workbook(
            reverse('reports:customer_report') + '?format=excel'
        )
        # Find the SUMMARY row and check the following rows
        summary_row = None
        for r in range(1, ws.max_row + 1):
            if ws.cell(r, 1).value == 'SUMMARY':
                summary_row = r
                break
        self.assertIsNotNone(summary_row)
        self.assertEqual(ws.cell(summary_row + 1, 1).value, 'Total Customers')
        self.assertEqual(ws.cell(summary_row + 2, 1).value, 'Active Customers')
        self.assertEqual(ws.cell(summary_row + 3, 1).value, 'Generated On')

    def test_customer_report_excel_has_table_columns(self):
        """All 15 required columns are present in the table header."""
        ws = self._load_report_workbook(
            reverse('reports:customer_report') + '?format=excel'
        )
        expected = [
            'S.N', 'Company Name', 'TRN', 'Trade License', 'TL Expiry',
            'TL Days Left', 'Passport', 'Passport Expiry', 'Passport Days Left',
            'EID', 'EID Expiry', 'EID Days Left', 'Copy Status',
            'Security Cheque', 'Sales Person'
        ]
        # Find header row: first row where col 1 == 'S.N'
        header_row = None
        for r in range(1, ws.max_row + 1):
            if ws.cell(r, 1).value == 'S.N':
                header_row = r
                break
        self.assertIsNotNone(header_row)
        for idx, col in enumerate(expected, 1):
            self.assertEqual(ws.cell(header_row, idx).value, col,
                             f'Column {idx} ({col}) mismatch')

    def test_customer_report_excel_has_auto_filter_and_freeze(self):
        """AutoFilter and freeze panes are configured on the table."""
        ws = self._load_report_workbook(
            reverse('reports:customer_report') + '?format=excel'
        )
        self.assertTrue(ws.auto_filter.ref)
        self.assertTrue(ws.freeze_panes)

    def test_customer_report_excel_empty_message(self):
        """An empty filtered result shows a friendly empty message."""
        ws = self._load_report_workbook(
            reverse('reports:customer_report')
            + '?format=excel&customer_status=inactive&search=__NO_SUCH_CUSTOMER__'
        )
        found = any(
            ws.cell(r, 1).value
            and 'No customer records found' in str(ws.cell(r, 1).value)
            for r in range(1, ws.max_row + 1)
        )
        self.assertTrue(found)

    def test_customer_report_excel_footer(self):
        """The footer contains the Notes and system branding."""
        ws = self._load_report_workbook(
            reverse('reports:customer_report') + '?format=excel'
        )
        found_notes = any(
            ws.cell(r, 1).value == 'NOTES' for r in range(1, ws.max_row + 1)
        )
        found_footer = any(
            ws.cell(r, 1).value == 'KB Remicon Customer Record System'
            for r in range(1, ws.max_row + 1)
        )
        self.assertTrue(found_notes, 'NOTES section missing')
        self.assertTrue(found_footer, 'Branding footer missing')

    def test_customer_report_excel_sales_person_subtitle(self):
        """A specific sales person appears as a subtitle under the title."""
        ws = self._load_report_workbook(
            reverse('reports:customer_report')
            + f'?format=excel&sales_person={self.sales_person.pk}'
        )
        self.assertEqual(ws['E3'].value, f'Sales Person: {self.sales_person.name}')

    def test_customer_report_excel_all_sales_persons_subtitle(self):
        """No specific sales person shows 'All Sales Persons' as subtitle."""
        ws = self._load_report_workbook(
            reverse('reports:customer_report') + '?format=excel'
        )
        self.assertEqual(ws['E3'].value, 'Sales Person: All Sales Persons')

    def test_customer_report_sales_person_filters_html(self):
        """HTML report only lists customers of the selected sales person."""
        other_sp = SalesPerson.objects.create(
            name='Jane Roe', email='jane@example.com'
        )
        Customer.objects.create(
            company_name='Other Company LLC',
            trade_license_number='99999',
            sales_person=other_sp,
            created_by=self.admin_user,
            updated_by=self.admin_user,
        )
        self.client.login(email='admin@example.com', password='admin123')
        response = self.client.get(
            reverse('reports:customer_report')
            + f'?sales_person={self.sales_person.pk}'
        )
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Test Company LLC')
        self.assertNotContains(response, 'Other Company LLC')

    def test_customer_report_sales_person_and_active_filter(self):
        """Sales Person + Active status returns only active customers of that
        sales person, excluding inactive ones even from the same person."""
        other_sp = SalesPerson.objects.create(
            name='Jane Roe', email='jane@example.com'
        )
        # Inactive customer of the same selected sales person.
        Customer.objects.create(
            company_name='Inactive Same SP',
            sales_person=self.sales_person,
            customer_status=Customer.CustomerStatus.INACTIVE,
            created_by=self.admin_user,
            updated_by=self.admin_user,
        )
        # Active customer of a different sales person.
        Customer.objects.create(
            company_name='Active Other SP',
            sales_person=other_sp,
            customer_status=Customer.CustomerStatus.ACTIVE,
            created_by=self.admin_user,
            updated_by=self.admin_user,
        )
        self.client.login(email='admin@example.com', password='admin123')
        response = self.client.get(
            reverse('reports:customer_report')
            + f'?sales_person={self.sales_person.pk}&customer_status=active'
        )
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Test Company LLC')
        self.assertNotContains(response, 'Inactive Same SP')
        self.assertNotContains(response, 'Active Other SP')

    def test_customer_report_excel_filename_has_sales_person(self):
        """Excel filename embeds the sanitized sales person name."""
        self.client.login(email='admin@example.com', password='admin123')
        response = self.client.get(
            reverse('reports:customer_report')
            + f'?format=excel&sales_person={self.sales_person.pk}'
        )
        self.assertEqual(response.status_code, 200)
        self.assertIn('filename=', response['Content-Disposition'])
        self.assertIn('John_Doe', response['Content-Disposition'])

    def test_customer_report_excel_filename_default(self):
        """Default Excel filename when no sales person selected."""
        self.client.login(email='admin@example.com', password='admin123')
        response = self.client.get(
            reverse('reports:customer_report') + '?format=excel'
        )
        self.assertEqual(response.status_code, 200)
        self.assertIn(
            'KB_Remicon_Customer_Compliance_Report.xlsx',
            response['Content-Disposition'],
        )

    def test_customer_report_pdf_filename_has_sales_person(self):
        """PDF filename embeds the sanitized sales person name."""
        self.client.login(email='admin@example.com', password='admin123')
        response = self.client.get(
            reverse('reports:customer_report')
            + f'?format=pdf&sales_person={self.sales_person.pk}'
        )
        self.assertEqual(response.status_code, 200)
        self.assertIn('John_Doe', response['Content-Disposition'])

    def test_customer_report_sales_person_zero_customers(self):
        """A sales person with no customers produces an empty report."""
        empty_sp = SalesPerson.objects.create(
            name='Empty SP', email='empty@example.com'
        )
        self.client.login(email='admin@example.com', password='admin123')
        response = self.client.get(
            reverse('reports:customer_report')
            + f'?sales_person={empty_sp.pk}'
        )
        self.assertEqual(response.status_code, 200)
        self.assertNotContains(response, 'Test Company LLC')

    def test_customer_report_sales_person_and_expired_filter(self):
        """Sales Person + expired trade-license filter."""
        other_sp = SalesPerson.objects.create(
            name='Jane Roe', email='jane@example.com'
        )
        # Expired customer of the selected sales person.
        Customer.objects.create(
            company_name='Expired Same SP',
            sales_person=self.sales_person,
            trade_license_expiry=date.today() - timedelta(days=10),
            created_by=self.admin_user,
            updated_by=self.admin_user,
        )
        # Expired customer of a different sales person (must be excluded).
        Customer.objects.create(
            company_name='Expired Other SP',
            sales_person=other_sp,
            trade_license_expiry=date.today() - timedelta(days=5),
            created_by=self.admin_user,
            updated_by=self.admin_user,
        )
        self.client.login(email='admin@example.com', password='admin123')
        response = self.client.get(
            reverse('reports:customer_report')
            + f'?sales_person={self.sales_person.pk}&trade_license_status=expired'
        )
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Expired Same SP')
        self.assertNotContains(response, 'Expired Other SP')

    def test_customer_report_sales_person_and_date_range(self):
        """Sales Person + date-from/date-to filters combine correctly."""
        # "Test Company LLC" was created in setUp today, so it falls within the
        # 30-day range and belongs to the selected sales person.
        older = Customer.objects.create(
            company_name='Older Same SP',
            sales_person=self.sales_person,
            created_by=self.admin_user,
            updated_by=self.admin_user,
        )
        # auto_now_add ignores a provided created_at, so set it explicitly after.
        Customer.objects.filter(pk=older.pk).update(
            created_at=timezone.now() - timedelta(days=400)
        )
        self.client.login(email='admin@example.com', password='admin123')
        # Date range covering only the last 30 days (excludes the 400-day-old one).
        response = self.client.get(
            reverse('reports:customer_report')
            + f'?sales_person={self.sales_person.pk}'
            + f'&date_from={(date.today() - timedelta(days=30)).isoformat()}'
            + f'&date_to={date.today().isoformat()}'
        )
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Test Company LLC')
        self.assertNotContains(response, 'Older Same SP')

    def test_excel_filename_sanitizes_sales_person_name(self):
        """Sales person names with illegal filename characters are sanitized."""
        from reports.excel_reports import _sanitize_filename_component
        self.assertEqual(_sanitize_filename_component('Mr. Moujahed'), 'Mr._Moujahed')
        self.assertEqual(_sanitize_filename_component('A/B:C*D?E"F<G>H|I'), 'A_B_C_D_E_F_G_H_I')

    def test_customer_report_all_sales_persons_shows_all(self):
        """All Sales Persons shows customers across every sales person."""
        other_sp = SalesPerson.objects.create(
            name='Jane Roe', email='jane@example.com'
        )
        Customer.objects.create(
            company_name='Other Company LLC',
            sales_person=other_sp,
            created_by=self.admin_user,
            updated_by=self.admin_user,
        )
        self.client.login(email='admin@example.com', password='admin123')
        response = self.client.get(reverse('reports:customer_report'))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Test Company LLC')
        self.assertContains(response, 'Other Company LLC')

    def test_reports_index_export_buttons_submit_form_filters(self):
        """Regression: the Reports Index export buttons must be form-submitting
        <button type=submit name=format> elements so the currently selected
        sales_person (and any other filter) is actually sent to the server.
        Previously they were <a href> links that only appended '?format=...' to
        the current URL, ignoring unsaved form values and exporting ALL data."""
        self.client.login(email='admin@example.com', password='admin123')
        response = self.client.get(reverse('reports:index'))
        self.assertEqual(response.status_code, 200)

        # Excel button must be a submit button carrying name/value 'format'.
        self.assertContains(
            response,
            '<button type="submit" name="format" value="excel"',
        )
        self.assertContains(
            response,
            '<button type="submit" name="format" value="pdf"',
        )
        # No export <a href> links should remain that bypass the form filters.
        self.assertNotContains(
            response,
            'href="{% url \'reports:customer_report\' %}?format=excel',
        )
        self.assertNotContains(
            response,
            'href="{% url \'reports:customer_report\' %}?format=pdf',
        )

    def test_customer_report_excel_filters_data_by_sales_person(self):
        """Exporting with a sales_person filter (as done via the Reports Index
        form) returns an Excel workbook containing ONLY that sales person's
        customers."""
        other_sp = SalesPerson.objects.create(
            name='Jane Roe', email='jane@example.com'
        )
        Customer.objects.create(
            company_name='Other Company LLC',
            trade_license_number='99999',
            sales_person=other_sp,
            created_by=self.admin_user,
            updated_by=self.admin_user,
        )
        self.client.login(email='admin@example.com', password='admin123')
        response = self.client.get(
            reverse('reports:customer_report')
            + f'?format=excel&sales_person={self.sales_person.pk}'
        )
        self.assertEqual(response.status_code, 200)
        ws = self._load_report_workbook(
            reverse('reports:customer_report')
            + f'?format=excel&sales_person={self.sales_person.pk}'
        )
        # Gather company names from the Company Name column (col 2).
        companies = []
        for r in range(1, ws.max_row + 1):
            val = ws.cell(r, 2).value
            if isinstance(val, str) and val not in ('Company Name',):
                companies.append(val)
        self.assertIn('Test Company LLC', companies)
        self.assertNotIn('Other Company LLC', companies)

    def test_customer_report_pdf_filters_data_by_sales_person(self):
        """Exporting with a sales_person filter (as done via the Reports Index
        form) returns a PDF whose filename references the sales person."""
        other_sp = SalesPerson.objects.create(
            name='Jane Roe', email='jane@example.com'
        )
        Customer.objects.create(
            company_name='Other Company LLC',
            trade_license_number='99999',
            sales_person=other_sp,
            created_by=self.admin_user,
            updated_by=self.admin_user,
        )
        self.client.login(email='admin@example.com', password='admin123')
        response = self.client.get(
            reverse('reports:customer_report')
            + f'?format=pdf&sales_person={self.sales_person.pk}'
        )
        self.assertEqual(response.status_code, 200)
        self.assertIn('John_Doe', response['Content-Disposition'])

    def test_missing_documents_sales_person_filter(self):
        """The Missing Documents report filters by the selected sales person."""
        other_sp = SalesPerson.objects.create(
            name='Jane Roe', email='jane@example.com'
        )
        missing_customer = Customer.objects.create(
            company_name='Missing Same SP',
            sales_person=self.sales_person,
            customer_status=Customer.CustomerStatus.ACTIVE,
            created_by=self.admin_user,
            updated_by=self.admin_user,
        )
        missing_customer.copy_status = Customer.CopyStatus.MISSING
        missing_customer.save()

        other_missing = Customer.objects.create(
            company_name='Missing Other SP',
            sales_person=other_sp,
            customer_status=Customer.CustomerStatus.ACTIVE,
            created_by=self.admin_user,
            updated_by=self.admin_user,
        )
        other_missing.copy_status = Customer.CopyStatus.MISSING
        other_missing.save()

        self.client.login(email='admin@example.com', password='admin123')
        response = self.client.get(
            reverse('reports:missing_documents')
            + f'?sales_person={self.sales_person.pk}'
        )
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Missing Same SP')
        self.assertNotContains(response, 'Missing Other SP')

    def test_missing_documents_report_lists_trn_certificate(self):
        """Regression: TRN Certificate is a required document, so the Missing
        Documents report must flag it for a customer missing only that copy."""
        customer = Customer.objects.create(
            company_name='Missing TRN Only LLC',
            sales_person=self.sales_person,
            created_by=self.admin_user,
            updated_by=self.admin_user,
        )
        # Has TL, Passport and EID copies but not a TRN Certificate.
        doc_args = {'customer': customer, 'uploaded_by': self.admin_user}
        Document.objects.create(
            document_type=Document.DocumentType.TRADE_LICENSE, file_name='tl.pdf',
            file_size=1, **doc_args,
        )
        Document.objects.create(
            document_type=Document.DocumentType.PASSPORT, file_name='pp.pdf',
            file_size=1, **doc_args,
        )
        Document.objects.create(
            document_type=Document.DocumentType.EMIRATES_ID, file_name='eid.pdf',
            file_size=1, **doc_args,
        )
        customer.copy_status = Customer.CopyStatus.PARTIAL
        customer.save(update_fields=['copy_status'])

        self.client.login(email='admin@example.com', password='admin123')
        response = self.client.get(reverse('reports:missing_documents'))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Missing TRN Only LLC')
        self.assertContains(response, 'TRN Certificate')

    def test_missing_documents_excel_reports_trn_certificate(self):
        """Regression: the missing-documents Excel report carries a TRN
        Certificate column and marks it 'No' for a customer lacking one."""
        customer = Customer.objects.create(
            company_name='Excel Missing TRN LLC',
            sales_person=self.sales_person,
            created_by=self.admin_user,
            updated_by=self.admin_user,
        )
        customer.copy_status = Customer.CopyStatus.MISSING
        customer.save(update_fields=['copy_status'])

        ws = self._load_report_workbook(
            reverse('reports:missing_documents') + '?format=excel'
        )
        header_row = None
        for r in range(1, ws.max_row + 1):
            if ws.cell(r, 1).value == 'S.N':
                header_row = r
                break
        self.assertIsNotNone(header_row)
        headers = [ws.cell(header_row, c).value for c in range(1, ws.max_column + 1)]
        self.assertIn('TRN Certificate', headers)
        trn_col = headers.index('TRN Certificate') + 1

        found = False
        for r in range(header_row + 1, ws.max_row + 1):
            if ws.cell(r, 2).value == 'Excel Missing TRN LLC':
                self.assertEqual(ws.cell(r, trn_col).value, 'No')
                found = True
                break
        self.assertTrue(found, 'customer row not found in report')

    def test_dashboard_and_notifications_pages_emit_no_bom(self):
        """Regression: rendered pages must not start with a UTF-8 BOM, which
        previously leaked from base.html / child templates into the browser."""
        self.client.login(email='admin@example.com', password='admin123')
        for url in ['/dashboard/', '/notifications/', '/customers/']:
            response = self.client.get(url)
            self.assertEqual(response.status_code, 200, url)
            content = response.content
            self.assertFalse(content.startswith(b'\xef\xbb\xbf'), url)
            if content:
                self.assertFalse(content.decode('utf-8', 'replace').startswith('\ufeff'), url)
    def test_expired_documents_sales_person_filter(self):
        """The Expired Documents report filters by the selected sales person."""
        other_sp = SalesPerson.objects.create(
            name='Jane Roe', email='jane@example.com'
        )
        Customer.objects.create(
            company_name='Expired Same SP',
            trade_license_expiry=date.today() - timedelta(days=10),
            sales_person=self.sales_person,
            created_by=self.admin_user,
            updated_by=self.admin_user,
        )
        Customer.objects.create(
            company_name='Expired Other SP',
            trade_license_expiry=date.today() - timedelta(days=10),
            sales_person=other_sp,
            created_by=self.admin_user,
            updated_by=self.admin_user,
        )
        self.client.login(email='admin@example.com', password='admin123')
        response = self.client.get(
            reverse('reports:expired_documents')
            + f'?sales_person={self.sales_person.pk}'
        )
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Expired Same SP')
        self.assertNotContains(response, 'Expired Other SP')

    def test_expiring_documents_sales_person_filter(self):
        """The Expiring Documents report filters by the selected sales person."""
        other_sp = SalesPerson.objects.create(
            name='Jane Roe', email='jane@example.com'
        )
        Customer.objects.create(
            company_name='Expiring Same SP',
            passport_expiry=date.today() + timedelta(days=20),
            sales_person=self.sales_person,
            created_by=self.admin_user,
            updated_by=self.admin_user,
        )
        Customer.objects.create(
            company_name='Expiring Other SP',
            passport_expiry=date.today() + timedelta(days=20),
            sales_person=other_sp,
            created_by=self.admin_user,
            updated_by=self.admin_user,
        )
        self.client.login(email='admin@example.com', password='admin123')
        response = self.client.get(
            reverse('reports:expiring_documents')
            + f'?sales_person={self.sales_person.pk}&days=30'
        )
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Expiring Same SP')
        self.assertNotContains(response, 'Expiring Other SP')

    def test_salesperson_report_customer_table(self):
        """Selecting a sales person shows their customer compliance table."""
        other_sp = SalesPerson.objects.create(
            name='Jane Roe', email='jane@example.com'
        )
        Customer.objects.create(
            company_name='Other SP Customer',
            sales_person=other_sp,
            created_by=self.admin_user,
            updated_by=self.admin_user,
        )
        self.client.login(email='admin@example.com', password='admin123')
        response = self.client.get(
            reverse('reports:salesperson_report')
            + f'?salesperson_id={self.sales_person.pk}'
        )
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Test Company LLC')
        self.assertNotContains(response, 'Other SP Customer')

    def test_salesperson_report_customer_table_excel(self):
        """The customer table mode exports a filtered customer Excel report."""
        self.client.login(email='admin@example.com', password='admin123')
        response = self.client.get(
            reverse('reports:salesperson_report')
            + f'?salesperson_id={self.sales_person.pk}&format=excel'
        )
        self.assertEqual(response.status_code, 200)
        self.assertIn(
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            response['Content-Type'],
        )





@pytest.mark.django_db
class TestProfileLogoUpload(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(
            email='profile@example.com',
            username='profileuser',
            password='testpass123',
            first_name='Profile',
            last_name='User',
            role=User.Role.ADMIN,
        )
        self.client.login(email='profile@example.com', password='testpass123')

    def _png_bytes(self):
        import io
        from PIL import Image
        buf = io.BytesIO()
        Image.new('RGBA', (10, 10), color=(0, 91, 172, 255)).save(buf, format='PNG')
        return buf.getvalue()

    def test_profile_edit_uploads_logo(self):
        import tempfile
        from django.test import override_settings
        from django.core.files.uploadedfile import SimpleUploadedFile

        image = SimpleUploadedFile(
            'logo.png', self._png_bytes(), content_type='image/png'
        )
        with tempfile.TemporaryDirectory() as tmp, override_settings(MEDIA_ROOT=tmp):
            response = self.client.post(
                reverse('accounts:profile_edit'),
                {
                    'first_name': 'Profile',
                    'last_name': 'User',
                    'phone': '+971 50 000 0000',
                    'email': 'profile@example.com',
                    'logo': image,
                },
            )
            # Successful edit redirects (302) to the profile detail page.
            self.assertEqual(response.status_code, 302)

            self.user.refresh_from_db()
            self.assertTrue(self.user.logo)
            self.assertIn('logos/', self.user.logo.name)

    def test_profile_edit_rejects_oversized_logo(self):
        import io
        import tempfile
        from django.test import override_settings
        from PIL import Image
        from django.core.files.uploadedfile import SimpleUploadedFile

        # Create a valid image but make it larger than the 5 MB limit.
        buf = io.BytesIO()
        Image.new('RGBA', (2560, 2560), color=(0, 91, 172, 255)).save(buf, format='PNG')
        png_bytes = buf.getvalue()
        # Pad the file to exceed 5 MB.
        oversized_bytes = png_bytes + b'\x00' * (5 * 1024 * 1024 + 1)
        big = SimpleUploadedFile(
            'big.png', oversized_bytes, content_type='image/png',
        )
        with tempfile.TemporaryDirectory() as tmp, override_settings(MEDIA_ROOT=tmp):
            response = self.client.post(
                reverse('accounts:profile_edit'),
                {
                    'first_name': 'Profile',
                    'last_name': 'User',
                    'phone': '+971 50 000 0000',
                    'email': 'profile@example.com',
                    'logo': big,
                },
            )
            self.assertEqual(response.status_code, 200)  # form re-rendered with error
            self.assertContains(response, 'too large')
            self.user.refresh_from_db()
            self.assertFalse(bool(self.user.logo))


@pytest.mark.django_db
class TestReportExcelExports(TestCase):
    def setUp(self):
        self.admin_user = User.objects.create_user(
            email='admin@example.com',
            username='admin',
            password='admin123',
            role=User.Role.ADMIN,
        )

    def test_salesperson_report_excel_export(self):
        """Regression: salesperson report Excel export must not raise a
        'Cannot convert ... to Excel' error."""
        self.client.login(email='admin@example.com', password='admin123')
        response = self.client.get(reverse('reports:salesperson_report') + '?format=excel')
        self.assertEqual(response.status_code, 200)
        self.assertIn(
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            response['Content-Type'],
        )

    def test_security_cheque_report_excel_export(self):
        """Regression: security cheque report Excel export must not raise a
        'Cannot convert ... to Excel' error."""
        self.client.login(email='admin@example.com', password='admin123')
        response = self.client.get(reverse('reports:cheque_report') + '?format=excel')
        self.assertEqual(response.status_code, 200)
        self.assertIn(
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            response['Content-Type'],
        )

@pytest.mark.django_db
class TestCustomerExcelImport(TestCase):
    """Regression tests for the Excel import flow (customers/import)."""

    def setUp(self):
        self.admin_user = User.objects.create_user(
            email='admin@example.com',
            username='admin',
            password='admin123',
            role=User.Role.ADMIN,
        )
        self.sales_person = SalesPerson.objects.create(
            name='Import Sales',
            email='import@example.com',
            phone='+971 50 000 0000',
        )

    def _make_xlsx(self, rows):
        """Return .xlsx bytes with the importer's expected headers and rows."""
        import io
        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        ws.append([
            'S.N', 'COMPANY NAME', 'TRADE LICENSE', 'EXPIRE DATE',
            'PASSPORT', 'PASSPORT EXPIRY', 'EID', 'EID EXPIRY',
            'TRN', 'SALES PERSON',
        ])
        for row in rows:
            ws.append(row)
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    def _upload(self, xlsx_bytes, filename='import.xlsx'):
        from django.core.files.uploadedfile import SimpleUploadedFile

        upload = SimpleUploadedFile(
            filename,
            xlsx_bytes,
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        return self.client.post(reverse('customers:import'), {'excel_file': upload})

    def _confirm_import(self):
        """Confirm the preview, follow the redirect to the status page, and
        poll the chunked importer until it reports done.

        The confirm request is intentionally fast: it only starts a pending
        job. The actual customer creation happens on the status page via
        chunked POST polls.
        """
        confirm = self.client.post(
            reverse('customers:import_preview'), {'confirm': '1'},
        )
        self.assertEqual(confirm.status_code, 302)
        self.assertEqual(confirm.url, reverse('customers:import_status'))

        status_page = self.client.get(reverse('customers:import_status'))
        self.assertEqual(status_page.status_code, 200)
        self.assertContains(status_page, 'progress-track')

        result = {}
        for _ in range(100):  # safe upper bound for the poll loop
            resp = self.client.post(reverse('customers:import_status'))
            self.assertEqual(resp.status_code, 200)
            result = resp.json()
            if result.get('done'):
                break
        self.assertTrue(result.get('done'), 'import never finished polling')
        return result

    def test_upload_preview_and_confirm_import(self):
        """Regression: string dates must parse, preview shows rows, confirm creates customers.
        Previously 'date.strptime' did not exist, so valid_rows was always empty."""
        self.client.login(email='admin@example.com', password='admin123')
        xlsx = self._make_xlsx([
            [1, 'Import Co A LLC', 'TL-A1', '2026-12-31',
             'PP-A1', '2026-06-30', 'EID-A1', '2026-09-30', 'TRN-A1', 'Import Sales'],
            [2, 'Import Co B LLC', 'TL-B1', None,
             'PP-B1', None, 'EID-B1', None, 'TRN-B1', ''],
        ])

        resp = self._upload(xlsx)
        self.assertEqual(resp.status_code, 302)
        self.assertEqual(resp.url, reverse('customers:import_preview'))

        # Preview page shows both valid rows and the record-count label
        preview = self.client.get(reverse('customers:import_preview'))
        self.assertEqual(preview.status_code, 200)
        self.assertContains(preview, 'Import Co A LLC')
        self.assertContains(preview, 'Import Co B LLC')
        self.assertContains(preview, 'Import 2 Records')

        # Confirm the import -- starts the pending job, then status polls it
        result = self._confirm_import()
        self.assertEqual(result['done'], True)
        self.assertIn('count', result)

        customer = Customer.objects.get(company_name='Import Co A LLC')
        self.assertEqual(customer.trade_license_expiry, date(2026, 12, 31))
        self.assertEqual(customer.passport_expiry, date(2026, 6, 30))
        self.assertEqual(customer.eid_expiry, date(2026, 9, 30))
        self.assertEqual(customer.trn, 'TRN-A1')
        self.assertEqual(customer.sales_person, self.sales_person)
        self.assertEqual(customer.created_by, self.admin_user)
        self.assertEqual(customer.customer_status, Customer.CustomerStatus.ACTIVE)

        customer_b = Customer.objects.get(company_name='Import Co B LLC')
        self.assertIsNone(customer_b.trade_license_expiry)
        self.assertIsNone(customer_b.sales_person)

    def test_import_with_excel_date_cells(self):
        """Regression: openpyxl yields datetime objects for real Excel date cells."""
        from datetime import datetime

        self.client.login(email='admin@example.com', password='admin123')
        xlsx = self._make_xlsx([
            [1, 'Import Co C LLC', 'TL-C1', datetime(2026, 12, 31),
             'PP-C1', datetime(2026, 6, 30), 'EID-C1', datetime(2026, 9, 30),
             'TRN-C1', 'Import Sales'],
        ])

        resp = self._upload(xlsx)
        self.assertEqual(resp.status_code, 302)

        preview = self.client.get(reverse('customers:import_preview'))
        self.assertContains(preview, 'Import Co C LLC')
        self.assertContains(preview, '2026-12-31')

        self._confirm_import()
        customer = Customer.objects.get(company_name='Import Co C LLC')
        self.assertEqual(customer.trade_license_expiry, date(2026, 12, 31))
        self.assertEqual(customer.passport_expiry, date(2026, 6, 30))

    def test_rejects_old_xls_files(self):
        """Only .xlsx is supported (openpyxl cannot read the legacy .xls format)."""
        self.client.login(email='admin@example.com', password='admin123')
        xlsx = self._make_xlsx([
            [1, 'Import Co D LLC', 'TL-D1', '2026-12-31',
             'PP-D1', '2026-06-30', 'EID-D1', '2026-09-30', 'TRN-D1', ''],
        ])
        resp = self._upload(xlsx, filename='customers.xls')
        self.assertEqual(resp.status_code, 200)  # form re-rendered with error
        self.assertContains(resp, 'File must be an Excel file (.xlsx).')
        self.assertFalse(
            Customer.objects.filter(company_name='Import Co D LLC').exists()
        )

    def test_duplicate_rows_are_reported_once(self):
        """A duplicate company in the DB is skipped and counted once."""
        self.client.login(email='admin@example.com', password='admin123')
        Customer.objects.create(
            company_name='Existing Co LLC',
            created_by=self.admin_user,
            updated_by=self.admin_user,
        )
        xlsx = self._make_xlsx([
            [1, 'Existing Co LLC', 'TL-E1', '2026-12-31',
             'PP-E1', '2026-06-30', 'EID-E1', '2026-09-30', 'TRN-E1', ''],
            [2, 'Fresh Co LLC', 'TL-E2', '2026-12-31',
             'PP-E2', '2026-06-30', 'EID-E2', '2026-09-30', 'TRN-E2', ''],
        ])

        resp = self._upload(xlsx)
        self.assertEqual(resp.status_code, 302)

        preview = self.client.get(reverse('customers:import_preview'))
        self.assertContains(preview, 'Fresh Co LLC')
        self.assertContains(preview, 'Duplicate company name')

        self._confirm_import()
        # Only the fresh company is created; the existing one is untouched.
        self.assertEqual(
            Customer.objects.filter(company_name='Fresh Co LLC').count(), 1
        )
        self.assertEqual(
            Customer.objects.filter(company_name='Existing Co LLC').count(), 1
        )
    def test_chunked_import_progress_polls(self):
        """Chunked import: the pending job is processed over several polls and
        intermediate responses report processed/total before done=True."""
        self.client.login(email='admin@example.com', password='admin123')
        rows = []
        for i in range(60):  # > CHUNK_SIZE (25), forces multiple polls
            rows.append([
                i + 1, 'Chunk Co {:02d} LLC'.format(i + 1), 'TL-CH{:02d}'.format(i + 1),
                '2026-12-31', 'PP-CH{:02d}'.format(i + 1), '2026-06-30',
                'EID-CH{:02d}'.format(i + 1), '2026-09-30', 'TRN-CH{:02d}'.format(i + 1),
                'Import Sales',
            ])
        resp = self._upload(self._make_xlsx(rows))
        self.assertEqual(resp.status_code, 302)

        self.client.post(reverse('customers:import_preview'), {'confirm': '1'})

        seen_partial = False
        total = None
        for _ in range(100):
            data = self.client.post(reverse('customers:import_status')).json()
            if data.get('done'):
                break
            seen_partial = True
            total = data.get('total')
            self.assertLessEqual(data['processed'], data['total'])
        self.assertTrue(seen_partial, 'expected at least one partial poll response')
        self.assertEqual(total, 60)
        self.assertEqual(Customer.objects.filter(company_name__startswith='Chunk Co').count(), 60)

    def test_cancel_import_aborts_pending_job(self):
        """Cancel returns done=True and creates no customers."""
        self.client.login(email='admin@example.com', password='admin123')
        xlsx = self._make_xlsx([
            [1, 'Cancelled Co LLC', 'TL-CAN1', '2026-12-31',
             'PP-CAN1', '2026-06-30', 'EID-CAN1', '2026-09-30', 'TRN-CAN1', ''],
        ])
        resp = self._upload(xlsx)
        self.assertEqual(resp.status_code, 302)

        self.client.post(reverse('customers:import_preview'), {'confirm': '1'})

        resp = self.client.post(reverse('customers:import_status'), {'cancel': '1'})
        self.assertEqual(resp.status_code, 200)
        data = resp.json()
        self.assertTrue(data['done'])
        self.assertEqual(data['redirect'], reverse('customers:import'))
        self.assertFalse(Customer.objects.filter(company_name='Cancelled Co LLC').exists())

    def test_status_without_pending_job(self):
        """Status page/API must be safe when no pending import exists."""
        self.client.login(email='admin@example.com', password='admin123')
        page = self.client.get(reverse('customers:import_status'))
        self.assertEqual(page.status_code, 302)
        self.assertEqual(page.url, reverse('customers:import'))

        data = self.client.post(reverse('customers:import_status')).json()
        self.assertTrue(data['done'])
        self.assertEqual(data['redirect'], reverse('customers:import'))
