"""
Views for customers app.
"""

import csv
from datetime import date, datetime
from django.contrib import messages
from django.contrib.auth.mixins import LoginRequiredMixin, UserPassesTestMixin
from django.db.models import Q, Count, Prefetch
from django.http import HttpResponse, JsonResponse
from django.shortcuts import render, get_object_or_404, redirect
from django.urls import reverse_lazy, reverse
from django.utils.translation import gettext_lazy as _
from django.views.generic import (
    ListView, DetailView, CreateView, UpdateView, DeleteView, View
)
from django.views.decorators.http import require_http_methods
from django.utils.decorators import method_decorator

from .models import Customer
from .forms import CustomerForm, CustomerFilterForm, CustomerImportForm
from .utils import get_days_left, get_expiry_status, get_expiry_filter_queryset
from documents.models import Document
from security_cheques.models import SecurityCheque
from salespersons.models import SalesPerson
from audit.utils import log_action


class CustomerAccessMixin(UserPassesTestMixin):
    """Mixin to check customer access permissions."""
    
    def test_func(self):
        customer = self.get_object()
        user = self.request.user
        
        if user.is_admin_user:
            return True
        
        if user.is_sales_person_user and hasattr(user, 'sales_person'):
            return customer.sales_person == user.sales_person
        
        return False
    
    def handle_no_permission(self):
        messages.error(self.request, _('You do not have permission to access this customer.'))
        return redirect('customers:list')


class CustomerListView(LoginRequiredMixin, ListView):
    model = Customer
    template_name = 'customers/list.html'
    context_object_name = 'customers'
    paginate_by = 25

    def get_template_names(self):
        """Render only the table partial for HTMX requests."""
        if getattr(self.request, 'htmx', None):
            return ['customers/partials/table.html']
        return ['customers/list.html']

    def get_queryset(self):
        queryset = Customer.objects.select_related('sales_person', 'created_by').prefetch_related(
            'documents', 'security_cheques'
        )
        
        # Apply user-based filtering
        if not self.request.user.is_admin_user:
            if hasattr(self.request.user, 'sales_person') and self.request.user.sales_person:
                queryset = queryset.filter(sales_person=self.request.user.sales_person)
            else:
                queryset = queryset.none()
        
        # Get filter parameters
        self.filter_form = CustomerFilterForm(self.request.GET)
        
        if self.filter_form.is_valid():
            # Search
            search = self.filter_form.cleaned_data.get('search')
            if search:
                queryset = queryset.filter(
                    Q(company_name__icontains=search) |
                    Q(trade_license_number__icontains=search) |
                    Q(passport_number__icontains=search) |
                    Q(eid_number__icontains=search) |
                    Q(trn__icontains=search) |
                    Q(sales_person__name__icontains=search)
                )
            
            # Sales person filter
            sales_person = self.filter_form.cleaned_data.get('sales_person')
            if sales_person:
                queryset = queryset.filter(sales_person=sales_person)
            
            # Customer status filter
            customer_status = self.filter_form.cleaned_data.get('customer_status')
            if customer_status:
                queryset = queryset.filter(customer_status=customer_status)
            
            # Document status filter
            document_status = self.filter_form.cleaned_data.get('document_status')
            if document_status:
                if document_status == 'expired':
                    queryset = queryset.filter(
                        Q(trade_license_expiry__lt=date.today()) |
                        Q(passport_expiry__lt=date.today()) |
                        Q(eid_expiry__lt=date.today())
                    )
                elif document_status == '7_days':
                    from datetime import timedelta
                    today = date.today()
                    week_later = today + timedelta(days=7)
                    queryset = queryset.filter(
                        Q(trade_license_expiry__gte=today, trade_license_expiry__lte=week_later) |
                        Q(passport_expiry__gte=today, passport_expiry__lte=week_later) |
                        Q(eid_expiry__gte=today, eid_expiry__lte=week_later)
                    )
                elif document_status == '30_days':
                    from datetime import timedelta
                    today = date.today()
                    month_later = today + timedelta(days=30)
                    queryset = queryset.filter(
                        Q(trade_license_expiry__gte=today, trade_license_expiry__lte=month_later) |
                        Q(passport_expiry__gte=today, passport_expiry__lte=month_later) |
                        Q(eid_expiry__gte=today, eid_expiry__lte=month_later)
                    )
                elif document_status == '60_days':
                    from datetime import timedelta
                    today = date.today()
                    two_months_later = today + timedelta(days=60)
                    queryset = queryset.filter(
                        Q(trade_license_expiry__gte=today, trade_license_expiry__lte=two_months_later) |
                        Q(passport_expiry__gte=today, passport_expiry__lte=two_months_later) |
                        Q(eid_expiry__gte=today, eid_expiry__lte=two_months_later)
                    )
                elif document_status == 'missing':
                    # Customers missing at least one required document
                    queryset = queryset.exclude(copy_status=Customer.CopyStatus.COMPLETE)
            
            # Copy status filter
            copy_status = self.filter_form.cleaned_data.get('copy_status')
            if copy_status:
                queryset = queryset.filter(copy_status=copy_status)
            
            # Cheque status filter
            cheque_status = self.filter_form.cleaned_data.get('cheque_status')
            if cheque_status:
                queryset = queryset.filter(security_cheques__status=cheque_status).distinct()
        
        # Sorting
        sort = self.request.GET.get('sort', '-created_at')
        allowed_sorts = [
            'company_name', '-company_name',
            'trade_license_number', '-trade_license_number',
            'trade_license_expiry', '-trade_license_expiry',
            'passport_expiry', '-passport_expiry',
            'eid_expiry', '-eid_expiry',
            'sales_person__name', '-sales_person__name',
            'customer_status', '-customer_status',
            'copy_status', '-copy_status',
            'created_at', '-created_at',
        ]
        if sort in allowed_sorts:
            queryset = queryset.order_by(sort)
        else:
            queryset = queryset.order_by('-created_at')
        
        return queryset.distinct()
    
    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['filter_form'] = self.filter_form
        context['sort'] = self.request.GET.get('sort', '-created_at')
        
        # Add computed fields for each customer
        for customer in context['customers']:
            customer.tl_days_left = customer.get_trade_license_days_left()
            customer.tl_status = customer.get_trade_license_status()
            customer.passport_days_left = customer.get_passport_days_left()
            customer.passport_status = customer.get_passport_status()
            customer.eid_days_left = customer.get_eid_days_left()
            customer.eid_status = customer.get_eid_status()
        
        return context


class CustomerDetailView(LoginRequiredMixin, CustomerAccessMixin, DetailView):
    model = Customer
    template_name = 'customers/detail.html'
    context_object_name = 'customer'
    
    def get_queryset(self):
        return Customer.objects.select_related('sales_person', 'created_by', 'updated_by').prefetch_related(
            'documents',
            'security_cheques',
        )
    
    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        customer = self.object
        
        # Add computed fields
        context['tl_days_left'] = customer.get_trade_license_days_left()
        context['tl_status'] = customer.get_trade_license_status()
        context['passport_days_left'] = customer.get_passport_days_left()
        context['passport_status'] = customer.get_passport_status()
        context['eid_days_left'] = customer.get_eid_days_left()
        context['eid_status'] = customer.get_eid_status()
        
        # Documents grouped by type
        context['documents_by_type'] = {}
        for doc in customer.documents.all():
            if doc.document_type not in context['documents_by_type']:
                context['documents_by_type'][doc.document_type] = []
            context['documents_by_type'][doc.document_type].append(doc)
        
        # Security cheques
        context['security_cheques'] = customer.security_cheques.all()
        
        return context


class CustomerCreateView(LoginRequiredMixin, CreateView):
    model = Customer
    form_class = CustomerForm
    template_name = 'customers/form.html'
    success_url = reverse_lazy('customers:list')
    
    def get_form_kwargs(self):
        kwargs = super().get_form_kwargs()
        kwargs['user'] = self.request.user
        return kwargs
    
    def form_valid(self, form):
        response = super().form_valid(form)
        log_action(
            user=self.request.user,
            action='created',
            customer=self.object,
            description=f'Created customer: {self.object.company_name}',
            request=self.request,
        )
        messages.success(self.request, _('Customer created successfully.'))
        
        if '_addanother' in self.request.POST:
            return redirect('customers:add')
        return response


class CustomerUpdateView(LoginRequiredMixin, CustomerAccessMixin, UpdateView):
    model = Customer
    form_class = CustomerForm
    template_name = 'customers/form.html'
    
    def get_form_kwargs(self):
        kwargs = super().get_form_kwargs()
        kwargs['user'] = self.request.user
        return kwargs
    
    def get_success_url(self):
        return reverse('customers:detail', kwargs={'pk': self.object.pk})
    
    def form_valid(self, form):
        response = super().form_valid(form)
        log_action(
            user=self.request.user,
            action='updated',
            customer=self.object,
            description=f'Updated customer: {self.object.company_name}',
            request=self.request,
        )
        messages.success(self.request, _('Customer updated successfully.'))
        return response


class CustomerDeleteView(LoginRequiredMixin, CustomerAccessMixin, DeleteView):
    model = Customer
    template_name = 'customers/confirm_delete.html'
    success_url = reverse_lazy('customers:list')
    
    def dispatch(self, request, *args, **kwargs):
        if not request.user.is_admin_user:
            messages.error(request, _('Only administrators can delete customers.'))
            return redirect('customers:list')
        return super().dispatch(request, *args, **kwargs)
    
    def delete(self, request, *args, **kwargs):
        customer = self.get_object()
        company_name = customer.company_name
        log_action(
            user=request.user,
            action='deleted',
            customer=customer,
            description=f'Deleted customer: {company_name}',
            request=request,
        )
        messages.success(request, _('Customer deleted successfully.'))
        return super().delete(request, *args, **kwargs)


class CustomerExportView(LoginRequiredMixin, View):
    """Export customers to Excel or CSV."""
    
    def get(self, request, *args, **kwargs):
        # Reuse the same filtering logic as list view
        list_view = CustomerListView()
        list_view.request = request
        queryset = list_view.get_queryset()
        
        export_format = request.GET.get('format', 'excel')
        
        if export_format == 'csv':
            return self.export_csv(queryset)
        else:
            return self.export_excel(queryset)
    
    def export_csv(self, queryset):
        response = HttpResponse(content_type='text/csv')
        response['Content-Disposition'] = f'attachment; filename="customers_{date.today()}.csv"'
        
        writer = csv.writer(response)
        writer.writerow([
            'S.N', 'Company Name', 'Trade License', 'Trade License Expiry', 'Trade License Days Left',
            'Passport', 'Passport Expiry', 'Passport Days Left',
            'EID', 'EID Expiry', 'EID Days Left',
            'Copy Status', 'TRN', 'Security Cheque', 'Sales Person'
        ])
        
        for i, customer in enumerate(queryset, 1):
            cheque = customer.security_cheques.first()
            writer.writerow([
                i,
                customer.company_name,
                customer.trade_license_number,
                customer.trade_license_expiry.strftime('%Y-%m-%d') if customer.trade_license_expiry else '',
                customer.get_trade_license_days_left() or '',
                customer.passport_number,
                customer.passport_expiry.strftime('%Y-%m-%d') if customer.passport_expiry else '',
                customer.get_passport_days_left() or '',
                customer.eid_number,
                customer.eid_expiry.strftime('%Y-%m-%d') if customer.eid_expiry else '',
                customer.get_eid_days_left() or '',
                customer.get_copy_status_display(),
                customer.trn,
                cheque.get_status_display() if cheque else 'Not Received',
                customer.sales_person.name if customer.sales_person else '',
            ])
        
        return response
    
    def export_excel(self, queryset):
        from reports.excel_reports import KBRemiconReport

        num_cols = 15
        headers = [
            'S.N', 'Company Name', 'Trade License', 'Trade License Expiry',
            'Trade License Days Left', 'Passport', 'Passport Expiry',
            'Passport Days Left', 'EID', 'EID Expiry', 'EID Days Left',
            'Copy Status', 'TRN', 'Security Cheque', 'Sales Person'
        ]

        report = KBRemiconReport('Customer List Export', num_cols=num_cols)
        report.add_header_block(
            'Customer List Export',
            subtitle=f'Total Customers: {queryset.count()}',
            report_type='Customer Record',
        )
        report.add_summary([
            ('Total Customers', str(queryset.count())),
            ('Generated', date.today().strftime('%d %b %Y')),
        ])
        report.add_table_headers(headers)

        for i, customer in enumerate(queryset, 1):
            cheque = customer.security_cheques.first()
            cheque_status = cheque.get_status_display() if cheque else 'Not Received'
            copy_status = customer.get_copy_status_display()
            report.add_data_row(
                [
                    i,
                    customer.company_name,
                    customer.trade_license_number or '',
                    customer.trade_license_expiry,
                    customer.get_trade_license_days_left(),
                    customer.passport_number or '',
                    customer.passport_expiry,
                    customer.get_passport_days_left(),
                    customer.eid_number or '',
                    customer.eid_expiry,
                    customer.get_eid_days_left(),
                    copy_status,
                    customer.trn or '',
                    cheque_status,
                    customer.sales_person.name if customer.sales_person else '',
                ],
                row_index=i,
                status_col=12,
                status_value=copy_status,
            )

        report.set_column_widths([
            6, 28, 16, 15, 15, 16, 15, 15, 16, 13, 13, 14, 16, 16, 20
        ])
        report.add_footer()
        report.setup_autofilter_freeze_print()

        return report.get_response(
            f'customers_{date.today()}.xlsx'
        )


class CustomerImportView(LoginRequiredMixin, UserPassesTestMixin, View):
    """Import customers from Excel."""
    
    def test_func(self):
        return self.request.user.is_admin_user
    
    def get(self, request):
        form = CustomerImportForm()
        return render(request, 'customers/import.html', {'form': form})
    
    def post(self, request):
        form = CustomerImportForm(request.POST, request.FILES)
        if form.is_valid():
            return self.process_import(request, form.cleaned_data['excel_file'])
        return render(request, 'customers/import.html', {'form': form})
    
    def process_import(self, request, excel_file):
        import openpyxl
        from io import BytesIO
        
        try:
            wb = openpyxl.load_workbook(BytesIO(excel_file.read()))
            ws = wb.active
        except Exception as e:
            messages.error(request, _('Error reading Excel file: {}').format(str(e)))
            return redirect('customers:import')
        
        # Normalize headers
        headers = []
        for cell in ws[1]:
            header = str(cell.value).strip().upper() if cell.value else ''
            headers.append(header)
        
        # Expected columns mapping
        column_map = {
            'S.N': 'sn',
            'COMPANY NAME': 'company_name',
            'COMPANY': 'company_name',
            'TRADE LICENSE': 'trade_license_number',
            'TRADE LICENCE': 'trade_license_number',
            'EXPIRE DATE': 'trade_license_expiry',
            'TRADE LICENSE EXPIRY': 'trade_license_expiry',
            'PASSPORT': 'passport_number',
            'PASSPORT EXPIRY': 'passport_expiry',
            'EID': 'eid_number',
            'EID EXPIRY': 'eid_expiry',
            'COPY': 'copy_status',
            'TRN': 'trn',
            'SECURITY CHEQUE': 'cheque_status',
            'SALES PERSON': 'sales_person',
        }
        
        # Map column indices
        col_indices = {}
        for idx, header in enumerate(headers):
            if header in column_map:
                col_indices[column_map[header]] = idx
        
        if 'company_name' not in col_indices:
            messages.error(request, _('Required column "Company Name" not found in Excel file.'))
            return redirect('customers:import')
        
        # Process rows
        valid_rows = []
        errors = []
        duplicates = 0
        row_count = 0
        seen_names = set()      # Track within-file duplicates
        seen_licenses = set()
        seen_trns = set()
        
        for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
            if not any(row):
                continue
            row_count += 1

            try:
                company_name = str(row[col_indices.get('company_name', 0)] or '').strip()
                if not company_name:
                    errors.append(f'Row {row_idx}: Company name is required.')
                    continue
                
                # Check for duplicate in DB and within this import file
                if Customer.objects.filter(company_name__iexact=company_name).exists():
                    duplicates += 1
                    errors.append(f'Row {row_idx}: Duplicate company name \"{company_name}\" (already in database).')
                    continue
                if company_name.lower() in seen_names:
                    duplicates += 1
                    errors.append(f'Row {row_idx}: Duplicate company name \"{company_name}\" (duplicated within this file).')
                    continue
                seen_names.add(company_name.lower())

                trade_license = str(row[col_indices.get('trade_license_number', 1)] or '').strip()
                if trade_license and Customer.objects.filter(trade_license_number__iexact=trade_license).exists():
                    duplicates += 1
                    errors.append(f'Row {row_idx}: Duplicate trade license \"{trade_license}\" (already in database).')
                    continue
                if trade_license and trade_license.lower() in seen_licenses:
                    duplicates += 1
                    errors.append(f'Row {row_idx}: Duplicate trade license \"{trade_license}\" (duplicated within this file).')
                    continue
                if trade_license:
                    seen_licenses.add(trade_license.lower())

                trn = str(row[col_indices.get('trn')] or '').strip() if 'trn' in col_indices else ''
                if trn and Customer.objects.filter(trn__iexact=trn).exists():
                    duplicates += 1
                    errors.append(f'Row {row_idx}: Duplicate TRN \"{trn}\" (already in database).')
                    continue
                if trn and trn.lower() in seen_trns:
                    duplicates += 1
                    errors.append(f'Row {row_idx}: Duplicate TRN \"{trn}\" (duplicated within this file).')
                    continue
                if trn:
                    seen_trns.add(trn.lower())

                # Parse dates
                def parse_date(val):
                    if not val:
                        return None
                    if isinstance(val, datetime):
                        return val.date()
                    if isinstance(val, date):
                        return val
                    if isinstance(val, str):
                        for fmt in ('%Y-%m-%d', '%d/%m/%Y', '%d-%m-%Y', '%Y/%m/%d'):
                            try:
                                return datetime.strptime(val.strip(), fmt).date()
                            except ValueError:
                                continue
                    return None
                
                trade_license_expiry = parse_date(row[col_indices.get('trade_license_expiry')] if 'trade_license_expiry' in col_indices else None)
                passport_expiry = parse_date(row[col_indices.get('passport_expiry')] if 'passport_expiry' in col_indices else None)
                eid_expiry = parse_date(row[col_indices.get('eid_expiry')] if 'eid_expiry' in col_indices else None)
                
                # Sales person
                sales_person = None
                sp_name = str(row[col_indices.get('sales_person')] or '').strip() if 'sales_person' in col_indices else ''
                if sp_name:
                    sales_person = SalesPerson.objects.filter(name__iexact=sp_name, active=True).first()
                    if not sales_person:
                        errors.append(f'Row {row_idx}: Sales person "{sp_name}" not found.')
                        # Continue without sales person
                
                valid_rows.append({
                    'company_name': company_name,
                    'trade_license_number': trade_license or None,
                    'trade_license_expiry': trade_license_expiry.isoformat() if trade_license_expiry else None,
                    'passport_number': str(row[col_indices.get('passport_number')] or '').strip() if 'passport_number' in col_indices else '',
                    'passport_expiry': passport_expiry.isoformat() if passport_expiry else None,
                    'eid_number': str(row[col_indices.get('eid_number')] or '').strip() if 'eid_number' in col_indices else '',
                    'eid_expiry': eid_expiry.isoformat() if eid_expiry else None,
                    'trn': trn or None,
                    'sales_person_pk': sales_person.pk if sales_person else None,
                    'sales_person_name': sales_person.name if sales_person else '',
                    'customer_status': Customer.CustomerStatus.ACTIVE,
                })
                
            except Exception as e:
                errors.append(f'Row {row_idx}: {str(e)}')
        
        # Store preview in session
        request.session['import_preview'] = {
            'valid_rows': valid_rows,
            'errors': errors,
            'duplicates': duplicates,
            'total_rows': row_count,
        }
        
        return redirect('customers:import_preview')


class CustomerImportPreviewView(LoginRequiredMixin, UserPassesTestMixin, View):
    """Show the import preview and start the chunked import."""
    
    def test_func(self):
        return self.request.user.is_admin_user
    
    def get(self, request):
        preview = request.session.get('import_preview')
        if not preview:
            messages.error(request, _('No import preview found. Please upload a file first.'))
            return redirect('customers:import')
        return render(request, 'customers/import_preview.html', {'preview': preview})
    
    def post(self, request):
        preview = request.session.get('import_preview')
        if not preview:
            messages.error(request, _('No import preview found.'))
            return redirect('customers:import')
        
        if 'confirm' not in request.POST:
            return redirect('customers:import')
        
        # Start the import: move the validated rows into a pending job the
        # status page processes in chunks. This keeps the confirm request
        # instant and shows the user real progress on the status page.
        request.session['import_pending'] = {
            'index': 0,
            'created': 0,
            'valid_rows': preview.get('valid_rows', []),
            'errors': preview.get('errors', []),
            'duplicates': preview.get('duplicates', 0),
        }
        del request.session['import_preview']
        request.session.modified = True
        return redirect('customers:import_status')


class CustomerImportStatusView(LoginRequiredMixin, UserPassesTestMixin, View):
    """Import customers from the pending job in chunks and report progress."""

    #: Rows imported per polling request.
    CHUNK_SIZE = 25
    
    def test_func(self):
        return self.request.user.is_admin_user
    
    def get(self, request):
        pending = request.session.get('import_pending')
        if not pending:
            messages.error(request, _('No import in progress. Please upload a file first.'))
            return redirect('customers:import')
        return render(request, 'customers/import_status.html', {
            'total': len(pending['valid_rows']),
            'processed': pending['index'],
        })
    
    def post(self, request):
        pending = request.session.get('import_pending')
        if not pending:
            return JsonResponse({'done': True, 'redirect': reverse('customers:import')})
        if 'cancel' in request.POST:
            self.clear_pending(request)
            messages.info(request, _('Import cancelled.'))
            return JsonResponse({'done': True, 'redirect': reverse('customers:import')})
        
        rows = pending['valid_rows']
        total = len(rows)
        start = pending['index']
        if start < total:
            end = min(start + self.CHUNK_SIZE, total)
            chunk_created, chunk_errors = self.import_rows(request, rows[start:end])
            pending['index'] = end
            pending['created'] += chunk_created
            pending['errors'].extend(chunk_errors)
            request.session.modified = True
        
        if pending['index'] >= total:
            created = pending.get('created', 0)
            errors = pending.get('errors', [])
            duplicates = pending.get('duplicates', 0)
            self.clear_pending(request)
            messages.success(request, _('Successfully imported {} customers.').format(created))
            if errors:
                messages.warning(request, _('{} rows had errors and were skipped.').format(len(errors)))
            if duplicates:
                messages.warning(request, _('{} duplicate rows were skipped.').format(duplicates))
            return JsonResponse({
                'done': True,
                'count': created,
                'redirect': reverse('customers:list'),
            })
        
        return JsonResponse({
            'done': False,
            'processed': pending['index'],
            'total': total,
        })
    def import_rows(self, request, rows):
        """Create customers for one chunk of rows. Returns (created, errors)."""
        pk_set = {r['sales_person_pk'] for r in rows if r.get('sales_person_pk')}
        sales_persons = {sp.pk: sp for sp in SalesPerson.objects.filter(pk__in=pk_set)}
        created = 0
        errors = []
        for row_data in rows:
            try:
                customer = Customer.objects.create(
                    company_name=row_data['company_name'],
                    trade_license_number=row_data.get('trade_license_number') or None,
                    trade_license_expiry=date.fromisoformat(row_data['trade_license_expiry']) if row_data.get('trade_license_expiry') else None,
                    passport_number=row_data.get('passport_number') or '',
                    passport_expiry=date.fromisoformat(row_data['passport_expiry']) if row_data.get('passport_expiry') else None,
                    eid_number=row_data.get('eid_number') or '',
                    eid_expiry=date.fromisoformat(row_data['eid_expiry']) if row_data.get('eid_expiry') else None,
                    trn=row_data.get('trn') or None,
                    sales_person=sales_persons.get(row_data.get('sales_person_pk')),
                    customer_status=row_data.get('customer_status', Customer.CustomerStatus.ACTIVE),
                    created_by=request.user,
                    updated_by=request.user,
                )
                customer.update_copy_status()
                log_action(
                    user=request.user,
                    action='imported',
                    customer=customer,
                    description='Imported customer: {}'.format(customer.company_name),
                    request=request,
                )
                created += 1
            except Exception as e:
                errors.append('Import error for {}: {}'.format(row_data['company_name'], str(e)))
        return created, errors
    
    def clear_pending(self, request):
        if 'import_pending' in request.session:
            del request.session['import_pending']
            request.session.modified = True

class CustomerPrintView(LoginRequiredMixin, CustomerAccessMixin, DetailView):
    """Print-friendly customer view."""
    model = Customer
    template_name = 'customers/print.html'
    context_object_name = 'customer'
    
    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        customer = self.object
        context['tl_days_left'] = customer.get_trade_license_days_left()
        context['tl_status'] = customer.get_trade_license_status()
        context['passport_days_left'] = customer.get_passport_days_left()
        context['passport_status'] = customer.get_passport_status()
        context['eid_days_left'] = customer.get_eid_days_left()
        context['eid_status'] = customer.get_eid_status()
        context['print_date'] = date.today()
        return context


# HTMX partial views
def customer_row_partial(request, pk):
    """Return a single customer row for HTMX updates."""
    customer = get_object_or_404(Customer, pk=pk)
    
    # Check permissions
    if not request.user.is_admin_user:
        if not (request.user.is_sales_person_user and 
                hasattr(request.user, 'sales_person') and 
                customer.sales_person == request.user.sales_person):
            return HttpResponse(status=403)
    
    tl_days_left = customer.get_trade_license_days_left()
    tl_status = customer.get_trade_license_status()
    passport_days_left = customer.get_passport_days_left()
    passport_status = customer.get_passport_status()
    eid_days_left = customer.get_eid_days_left()
    eid_status = customer.get_eid_status()
    cheque = customer.security_cheques.first()
    
    return render(request, 'customers/partials/customer_row.html', {
        'customer': customer,
        'tl_days_left': tl_days_left,
        'tl_status': tl_status,
        'passport_days_left': passport_days_left,
        'passport_status': passport_status,
        'eid_days_left': eid_days_left,
        'eid_status': eid_status,
        'cheque': cheque,
    })